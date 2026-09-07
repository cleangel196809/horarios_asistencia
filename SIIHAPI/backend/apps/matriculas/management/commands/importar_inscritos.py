"""
SIIHAPI - Importa el reporte "INSCRITOS POR CICLO" (Excel institucional)
a Estudiante/Programa/Materia/Matricula, para un ciclo de formacion dado.

Uso:
    python manage.py importar_inscritos ruta/al/archivo.xlsx --periodo 2026-3T
    python manage.py importar_inscritos ruta/al/archivo.xlsx --periodo 2026-3T --dry-run

Requisitos previos:
    - El ciclo (--periodo, por defecto 2026-3T) YA debe existir en el
      sistema. Créalo primero con el botón "Nuevo Ciclo de Formación" del
      dashboard de Admin/Secretaría Académica/Decano (usa las fechas reales
      del ciclo), o con Periodo.objects.create(...) en el shell.

Qué hace, por cada fila del Excel:
    1. Busca/crea el Usuario del estudiante por cédula (IDENTIFICACION).
       Las cuentas nuevas quedan con clave inutilizable (no pueden loguear
       hasta que alguien les asigne una contraseña) — este import es para
       datos académicos, no para dar de alta accesos.
    2. Busca/crea su perfil de Estudiante (solo al crear: no se pisa el
       programa/código de un estudiante que ya existía).
    3. Busca/crea el Programa (por nombre) y la Materia (por programa +
       código + plan + ciclo) si hiciera falta.
    4. Crea o actualiza la Matrícula (estudiante, materia, periodo).
       ESTADO del Excel -> Matricula.estado:
           MATRICULADO -> ACTIVA
           INSCRITO    -> INSCRITA   (pre-registro, aún no confirmado)
       Una matrícula existente nunca se "degrada": si ya estaba ACTIVA y el
       Excel trae INSCRITO para la misma fila, se deja ACTIVA.

Es IDEMPOTENTE: se puede correr más de una vez con el mismo archivo (o uno
corregido) sin duplicar nada -- todo se resuelve por claves naturales
(cédula, código de materia, código de periodo).

Fase 3 (2026-09-04).
"""
import re
import unicodedata

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.academico.models import Facultad, Materia, Programa
from apps.autenticacion.models import Usuario
from apps.matriculas.models import Estudiante, Matricula, Periodo

COLUMNAS_REQUERIDAS = [
    'CICLO_INGRESO', 'TIPO', 'IDENTIFICACION', 'NOMBRES', 'APELLIDOS',
    'EMAIL', 'COD_PLAN', 'NOM_PLAN', 'COD_ASIGNATURA', 'ASIGNATURA',
    'CICLO', 'CREDITOS', 'ESTADO', 'NOMBRE_FACULTAD',
]

ESTADO_MAP = {
    'MATRICULADO': 'ACTIVA',
    'INSCRITO': 'INSCRITA',
}

# Nunca "bajar" una matrícula ya confirmada a un estado más débil si el
# Excel trae una fila más floja para la misma (estudiante, materia, periodo).
FUERZA_ESTADO = {'INSCRITA': 0, 'ACTIVA': 1, 'FINALIZADA': 2, 'CANCELADA': 2}


def _texto(v):
    if v is None:
        return ''
    return str(v).strip()


def _entero_como_texto(v):
    """openpyxl a veces trae cédulas/celulares como float (1000001690.0)."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _norm(s):
    """Mayúsculas + sin tildes, para comparar nombres de facultad/programa
    sin que un acento distinto haga fallar el match (dato real de Excel)."""
    s = unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().upper()


def _decimal(v, default):
    try:
        if v is None or v == '':
            return default
        return float(v)
    except (TypeError, ValueError):
        return default


def _slug_codigo(nombre, maxlen=20):
    """Codigo de Programa auto-generado cuando no hay uno en el Excel
    (el Excel trae COD_PLAN, que es el plan de estudios de la MATERIA, no
    el código del programa) -- se revisa a mano si hace falta."""
    n = unicodedata.normalize('NFKD', nombre).encode('ascii', 'ignore').decode()
    n = re.sub(r'[^A-Za-z0-9]+', '-', n).strip('-').upper()
    return (n[:maxlen] or 'PROG').rstrip('-')


def _tipo_programa(nombre):
    n = nombre.upper()
    if 'TECNOLOG' in n:
        return 'TEC'
    if 'TECNICO' in n or 'TÉCNICO' in n:
        return 'TL'
    if 'INGLES' in n or 'INGLÉS' in n:
        return 'ING'
    return 'PROF'


class Command(BaseCommand):
    help = 'Importa el Excel "INSCRITOS POR CICLO" a Estudiante/Matricula para un ciclo dado.'

    def add_arguments(self, parser):
        parser.add_argument('archivo', help='Ruta al .xlsx de inscritos')
        parser.add_argument('--periodo', default='2026-3T', help='Código del ciclo (debe existir ya). Default: 2026-3T')
        parser.add_argument('--dry-run', action='store_true', help='Solo muestra el resumen, no guarda nada')

    def handle(self, *args, **opts):
        ruta = opts['archivo']
        cod_periodo = opts['periodo'].strip().upper()
        dry_run = opts['dry_run']

        try:
            periodo = Periodo.objects.get(codigo=cod_periodo)
        except Periodo.DoesNotExist:
            raise CommandError(
                f'El ciclo "{cod_periodo}" no existe todavía. Créalo primero desde el '
                f'botón "Nuevo Ciclo de Formación" (Admin/Secretaría Académica/Decano) '
                f'con sus fechas reales, y vuelve a correr este import.'
            )

        wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        ws = wb.worksheets[0]

        header_row_idx = None
        headers = {}
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
            valores = [_texto(c) for c in row]
            if 'CICLO_INGRESO' in valores and 'IDENTIFICACION' in valores:
                header_row_idx = i
                headers = {v: idx for idx, v in enumerate(valores) if v}
                break
        if header_row_idx is None:
            raise CommandError('No encontré la fila de encabezados (busco CICLO_INGRESO / IDENTIFICACION).')

        faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in headers]
        if faltantes:
            raise CommandError(f'Faltan columnas esperadas en el Excel: {", ".join(faltantes)}')

        def col(row, nombre):
            idx = headers.get(nombre)
            return row[idx] if idx is not None and idx < len(row) else None

        stats = {
            'filas': 0, 'saltadas_ciclo_distinto': 0, 'saltadas_sin_datos': 0,
            'usuarios_creados': 0, 'estudiantes_creados': 0,
            'programas_creados': 0, 'materias_creadas': 0,
            'facultades_no_encontradas': set(),
            'matriculas_creadas': 0, 'matriculas_actualizadas': 0, 'matriculas_sin_cambio': 0,
        }
        programas_para_revisar = []

        facultad_cache = {_norm(f.nombre): f for f in Facultad.objects.all()}
        programa_cache = {_norm(p.nombre): p for p in Programa.objects.select_related('facultad')}
        materia_cache = {}
        for m in Materia.objects.all():
            materia_cache[(m.programa_id, m.codigo, m.plan or '', m.ciclo or '')] = m
        usuario_cache = {u.cedula: u for u in Usuario.objects.filter(cedula__isnull=False)}
        estudiante_cache = {}
        for e in Estudiante.objects.all():
            estudiante_cache[e.usuario_id] = e

        def get_or_create_programa(nombre, cod_plan, nombre_facultad):
            key = _norm(nombre)
            if key in programa_cache:
                return programa_cache[key]
            facultad = facultad_cache.get(_norm(nombre_facultad))
            if not facultad:
                stats['facultades_no_encontradas'].add(nombre_facultad)
                return None
            codigo = _slug_codigo(nombre)
            base_codigo, n = codigo, 1
            while Programa.objects.filter(codigo=codigo).exists():
                n += 1
                codigo = f'{base_codigo[:17]}-{n}'
            programa = Programa.objects.create(
                facultad=facultad, codigo=codigo, nombre=nombre.strip(),
                tipo=_tipo_programa(nombre), modalidad='PRES',
            )
            programa_cache[key] = programa
            stats['programas_creados'] += 1
            programas_para_revisar.append(f'{codigo} · {nombre.strip()}')
            return programa

        def get_or_create_materia(programa, cod_asig, nombre_asig, plan, ciclo, creditos):
            key = (programa.id_programa, cod_asig, plan, ciclo)
            if key in materia_cache:
                return materia_cache[key]
            materia = Materia.objects.create(
                programa=programa, codigo=cod_asig, nombre=nombre_asig or cod_asig,
                plan=plan or None, ciclo=ciclo or '1', creditos=_decimal(creditos, 2),
            )
            materia_cache[key] = materia
            stats['materias_creadas'] += 1
            return materia

        def get_or_create_estudiante(usuario, programa):
            if usuario.pk in estudiante_cache:
                return estudiante_cache[usuario.pk]
            estudiante = Estudiante.objects.create(
                usuario=usuario, codigo=usuario.cedula, programa=programa,
            )
            estudiante_cache[usuario.pk] = estudiante
            stats['estudiantes_creados'] += 1
            return estudiante

        def get_or_create_usuario(cedula, nombres, apellidos, email, telefono):
            if cedula in usuario_cache:
                return usuario_cache[cedula]
            correo = email.strip().lower()
            usuario = Usuario.objects.filter(correo=correo).first()
            if usuario:
                usuario_cache[cedula] = usuario
                return usuario
            usuario = Usuario(
                correo=correo, nombre=nombres.strip() or 'Estudiante',
                apellido=apellidos.strip(), cedula=cedula, rol='ESTUDIANTE',
                telefono=telefono,
            )
            usuario.set_unusable_password()
            usuario.save()
            usuario_cache[cedula] = usuario
            stats['usuarios_creados'] += 1
            return usuario

        rows = ws.iter_rows(min_row=header_row_idx + 1, values_only=True)

        def procesar():
            for row in rows:
                cedula = _entero_como_texto(col(row, 'IDENTIFICACION'))
                cod_asig = _texto(col(row, 'COD_ASIGNATURA'))
                if not cedula or not cod_asig:
                    stats['saltadas_sin_datos'] += 1
                    continue
                stats['filas'] += 1

                nombres = _texto(col(row, 'NOMBRES'))
                apellidos = _texto(col(row, 'APELLIDOS'))
                email = _texto(col(row, 'EMAIL')) or _texto(col(row, 'CORREO AULA VIRTUAL'))
                telefono = _entero_como_texto(col(row, 'CELULAR')) or _entero_como_texto(col(row, 'TELEFONO'))
                nom_plan = _texto(col(row, 'NOM_PLAN'))
                cod_plan = _texto(col(row, 'COD_PLAN'))
                nom_asig = _texto(col(row, 'ASIGNATURA'))
                ciclo_materia = _entero_como_texto(col(row, 'CICLO')) or '1'
                creditos = col(row, 'CREDITOS')
                estado_excel = _texto(col(row, 'ESTADO')).upper()
                facultad_nombre = _texto(col(row, 'NOMBRE_FACULTAD'))

                if not email or not nom_plan or not cod_asig:
                    stats['saltadas_sin_datos'] += 1
                    continue

                programa = get_or_create_programa(nom_plan, cod_plan, facultad_nombre)
                if not programa:
                    stats['saltadas_sin_datos'] += 1
                    continue

                usuario = get_or_create_usuario(cedula, nombres, apellidos, email, telefono)
                estudiante = get_or_create_estudiante(usuario, programa)
                materia = get_or_create_materia(programa, cod_asig, nom_asig, cod_plan, ciclo_materia, creditos)

                nuevo_estado = ESTADO_MAP.get(estado_excel, 'INSCRITA')
                matricula, creada = Matricula.objects.get_or_create(
                    estudiante=estudiante, materia=materia, periodo=periodo,
                    defaults={'estado': nuevo_estado},
                )
                if creada:
                    stats['matriculas_creadas'] += 1
                elif FUERZA_ESTADO.get(nuevo_estado, 0) > FUERZA_ESTADO.get(matricula.estado, 0):
                    matricula.estado = nuevo_estado
                    matricula.save(update_fields=['estado'])
                    stats['matriculas_actualizadas'] += 1
                else:
                    stats['matriculas_sin_cambio'] += 1

        if dry_run:
            with transaction.atomic():
                procesar()
                transaction.set_rollback(True)
        else:
            with transaction.atomic():
                procesar()

        self.stdout.write(self.style.SUCCESS(
            f'\n{"[DRY-RUN] " if dry_run else ""}Import de "{cod_periodo}" completo:'
        ))
        self.stdout.write(f'  Filas procesadas:            {stats["filas"]}')
        self.stdout.write(f'  Filas saltadas (sin datos):  {stats["saltadas_sin_datos"]}')
        self.stdout.write(f'  Usuarios (estudiantes) nuevos: {stats["usuarios_creados"]}')
        self.stdout.write(f'  Perfiles de Estudiante nuevos: {stats["estudiantes_creados"]}')
        self.stdout.write(f'  Programas nuevos:            {stats["programas_creados"]}')
        self.stdout.write(f'  Materias nuevas:             {stats["materias_creadas"]}')
        self.stdout.write(f'  Matrículas creadas:          {stats["matriculas_creadas"]}')
        self.stdout.write(f'  Matrículas actualizadas:     {stats["matriculas_actualizadas"]}')
        self.stdout.write(f'  Matrículas sin cambio:       {stats["matriculas_sin_cambio"]}')
        if programas_para_revisar:
            self.stdout.write(self.style.WARNING(
                '\n  Programas creados automáticamente (revisar tipo/código a mano):'))
            for p in programas_para_revisar:
                self.stdout.write(f'    - {p}')
        if stats['facultades_no_encontradas']:
            self.stdout.write(self.style.WARNING(
                '\n  Facultades del Excel que NO existen en el sistema (filas saltadas):'))
            for f in stats['facultades_no_encontradas']:
                self.stdout.write(f'    - {f}')
        if stats['usuarios_creados']:
            self.stdout.write(self.style.WARNING(
                f'\n  Los {stats["usuarios_creados"]} usuarios nuevos quedaron con clave '
                f'inutilizable (no pueden iniciar sesión hasta que se les asigne contraseña).'
            ))
