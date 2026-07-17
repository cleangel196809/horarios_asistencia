"""
SISCA — Generador de QR y exportación de reportes (PDF / Excel)
"""
import os
import uuid
import tempfile
from datetime import datetime, timedelta

import qrcode


# ── QR ──────────────────────────────────────────────────────────────────────

def crear_qr(id_sesion: int, ip: str, port, expiry_min: int) -> tuple:
    """
    Genera un token único, crea la imagen QR y devuelve
    (token, ruta_imagen, datetime_expiracion).
    """
    token = uuid.uuid4().hex
    url = f"http://{ip}:{port}/asistencia/registrar/{token}"
    fecha_exp = datetime.now() + timedelta(minutes=expiry_min)

    img = qrcode.make(url)
    path = os.path.join(tempfile.gettempdir(), f"sisca_qr_{token}.png")
    img.save(path)

    return token, path, fecha_exp


# ── PDF ─────────────────────────────────────────────────────────────────────

def generar_pdf_asistencia(datos: list) -> str:
    """Genera PDF de asistencia con ReportLab y devuelve la ruta del archivo."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Table,
                                    TableStyle, Paragraph, Spacer)
    from reportlab.lib.styles import getSampleStyleSheet

    path = os.path.join(tempfile.gettempdir(),
                        f"sisca_reporte_{uuid.uuid4().hex[:8]}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4)
    styles = getSampleStyleSheet()

    elements = [
        Paragraph("SISCA · Politécnico Internacional", styles["Title"]),
        Paragraph("Reporte de Asistencia", styles["Heading2"]),
        Paragraph(
            f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]

    headers = ["Estudiante", "Materia", "Fecha", "Estado", "Tipo"]
    rows = [headers] + [
        [
            str(d.get("estudiante", "")),
            str(d.get("nombre_materia", "")),
            str(d.get("fecha_sesion", ""))[:10],
            str(d.get("estado", "")),
            str(d.get("tipo_registro", "")),
        ]
        for d in datos
    ]

    tbl = Table(rows, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F6FEB")),
        ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",   (0, 0), (-1, -1), 9),
        ("GRID",       (0, 0), (-1, -1), 0.4, colors.HexColor("#30363D")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F6F8FA")]),
    ]))
    elements.append(tbl)
    doc.build(elements)
    return path


# ── Excel ────────────────────────────────────────────────────────────────────

def generar_excel_asistencia(datos: list) -> str:
    """Genera Excel de asistencia con openpyxl y devuelve la ruta."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asistencia SISCA"

    headers = ["Estudiante", "Materia", "Fecha Sesión",
               "Estado", "Tipo Registro", "Hora Registro"]
    hdr_fill = PatternFill("solid", fgColor="1F6FEB")
    hdr_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[chr(64 + col)].width = 22

    alt_fill = PatternFill("solid", fgColor="F6F8FA")
    for row_idx, d in enumerate(datos, 2):
        ws.append([
            d.get("estudiante", ""),
            d.get("nombre_materia", ""),
            str(d.get("fecha_sesion", ""))[:10],
            d.get("estado", ""),
            d.get("tipo_registro", ""),
            str(d.get("hora_registro", ""))[:19],
        ])
        if row_idx % 2 == 0:
            for col in range(1, 7):
                ws.cell(row=row_idx, column=col).fill = alt_fill

    path = os.path.join(tempfile.gettempdir(),
                        f"sisca_{uuid.uuid4().hex[:8]}.xlsx")
    wb.save(path)
    return path
